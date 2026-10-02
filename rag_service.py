# RAG Pipeline from scratch for better understanding 
# Developer - Amey Khodke (Optimized & Multilingual Enhanced)
# Decoupled Core Architecture for FastAPI Backend + Streamlit Dashboard
import io
import os
import re
import sys
import json
import time
import hashlib
import functools
import threading
from pathlib import Path
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor
from dotenv import load_dotenv
import chromadb
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import pypdf
import docx
from sentence_transformers import SentenceTransformer
from huggingface_hub.utils import disable_progress_bars

disable_progress_bars()

# Load environment variables cleanly
load_dotenv(Path(__file__).parent / ".env")
load_dotenv()

# Safe no-op decorator for backwards compatibility
def cache_resource(func=None, **kwargs):
    if func is None:
        return lambda f: f
    return func

# Module-level precompiled regex patterns
RE_LINE_BREAKS = re.compile(r'\r\n')
RE_WORD_TOKENS = re.compile(r'\w+', re.IGNORECASE)
RE_MULTIPLE_NEWLINES = re.compile(r'[\r\n]+')
RE_MULTIPLE_SPACES = re.compile(r'\s+')
RE_PAGE_NUMBERS = re.compile(r'\bPage\s*\d+\b', re.IGNORECASE)
RE_PP_NUMBERS = re.compile(r'\bpp\.\s*\d+\b', re.IGNORECASE)
RE_LONG_WORDS = re.compile(r'\b\w{5,}\b', re.IGNORECASE)

# Sentence boundary regex: respects decimals (3.14), abbreviations, and Devanagari danda (। / ॥)
ABBREVIATIONS = {
    'mr.', 'mrs.', 'ms.', 'dr.', 'prof.', 'sr.', 'jr.', 'vs.', 'etc.', 'e.g.',
    'i.e.', 'approx.', 'fig.', 'al.', 'dept.', 'inc.', 'ltd.', 'corp.', 'u.s.',
    'jan.', 'feb.', 'mar.', 'apr.', 'aug.', 'sept.', 'oct.', 'nov.', 'dec.'
}
PUNCT_SPLIT = re.compile(r'(\n+|[!?\u0964\u0965]+|(?<!\d)\.(?!\d))(?:\s+|$)')

# Reuse a single requests.Session() with connection pooling and retries
def _create_http_session():
    session = requests.Session()
    retries = Retry(
        total=2,
        backoff_factor=0.3,
        status_forcelist=[429, 500, 502, 503, 504],
        raise_on_status=False
    )
    adapter = HTTPAdapter(max_retries=retries)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session

HTTP_SESSION = _create_http_session()

# Cache directory & Schema version
CACHE_DIR = Path(__file__).parent / ".cache"
CACHE_DIR.mkdir(exist_ok=True)
SCHEMA_VERSION = 2
REGISTRY_FILE = CACHE_DIR / "registry.json"

# Thread-safe LRU Query Cache keyed on (normalized_question, language, tuple(sorted(doc_ids)))
QUERY_CACHE: OrderedDict = OrderedDict()
QUERY_CACHE_LOCK = threading.Lock()
MAX_QUERY_CACHE_SIZE = 500

# Language Mapping Dict (Extensible for 4th/5th languages)
LANGUAGE_CODES = {
    "English": "en-IN",
    "हिंदी (Hindi)": "hi-IN",
    "मराठी (Marathi)": "mr-IN"
}

def _get_config_value(keys: list[str]) -> str | None:
    """Read config from environment, or from st.secrets only if streamlit is already loaded."""
    for key in keys:
        val = os.getenv(key)
        if val:
            return val
    if "streamlit" in sys.modules:
        try:
            st = sys.modules["streamlit"]
            if hasattr(st, "secrets") and st.secrets:
                for key in keys:
                    if key in st.secrets and st.secrets[key]:
                        return st.secrets[key]
                for sec_name in ["groq", "sarvam"]:
                    sec = st.secrets.get(sec_name)
                    if isinstance(sec, dict):
                        for k in ["api_key", "API_KEY"]:
                            if k in sec and sec[k]:
                                return sec[k]
                    elif hasattr(sec, "get"):
                        for k in ["api_key", "API_KEY"]:
                            v = sec.get(k)
                            if v:
                                return v
        except Exception:
            pass
    return None

def get_groq_api_key():
    return _get_config_value(["GROQ_API_KEY", "groq_api_key"])

def get_sarvam_api_key():
    return _get_config_value(["SARVAM_API_KEY", "sarvam_api_key"])

def get_hf_token():
    return _get_config_value(["HF_TOKEN", "hf_token"])

# Singleton Model & ChromaDB Client
@functools.lru_cache(maxsize=1)
def get_embedding_model():
    print("Loading embedding model (singleton)...")
    return SentenceTransformer("all-MiniLM-L6-v2")

@functools.lru_cache(maxsize=1)
def get_chroma_client():
    print("Initializing ChromaDB Client (singleton)...")
    return chromadb.Client()

# Document Registry Management (Persisted to disk)
def load_document_registry() -> dict:
    if REGISTRY_FILE.exists():
        try:
            return json.loads(REGISTRY_FILE.read_text(encoding="utf-8"))
        except Exception as e:
            print(f"Error loading registry: {e}")
    return {}

def save_document_registry(registry: dict):
    try:
        REGISTRY_FILE.write_text(json.dumps(registry, indent=2), encoding="utf-8")
    except Exception as e:
        print(f"Error saving registry: {e}")

def list_registered_documents() -> list[dict]:
    registry = load_document_registry()
    return list(registry.values())

def get_registered_document(doc_id: str) -> dict | None:
    registry = load_document_registry()
    return registry.get(doc_id)

def delete_registered_document(doc_id: str) -> bool:
    client = get_chroma_client()
    try:
        collection = client.get_or_create_collection("document_embeddings")
        collection.delete(where={"doc_id": doc_id})
    except Exception as e:
        print(f"Error deleting doc {doc_id} from Chroma: {e}")

    registry = load_document_registry()
    doc_info = registry.pop(doc_id, None)
    save_document_registry(registry)
    clear_query_cache(doc_id=doc_id)
    return doc_info is not None

# Query Cache Helpers
def get_query_cache_entry(key: tuple):
    with QUERY_CACHE_LOCK:
        if key in QUERY_CACHE:
            QUERY_CACHE.move_to_end(key)
            return QUERY_CACHE[key]
    return None

def set_query_cache_entry(key: tuple, value):
    with QUERY_CACHE_LOCK:
        QUERY_CACHE[key] = value
        QUERY_CACHE.move_to_end(key)
        if len(QUERY_CACHE) > MAX_QUERY_CACHE_SIZE:
            QUERY_CACHE.popitem(last=False)

def clear_query_cache(doc_id: str | None = None):
    with QUERY_CACHE_LOCK:
        if doc_id is None:
            QUERY_CACHE.clear()
        else:
            keys_to_del = [k for k in QUERY_CACHE.keys() if doc_id in k[2]]
            for k in keys_to_del:
                del QUERY_CACHE[k]

# Document Loading & Extraction
def _extract_pdf_page(args):
    content_bytes, page_idx = args
    try:
        reader = pypdf.PdfReader(io.BytesIO(content_bytes))
        if page_idx < len(reader.pages):
            text = reader.pages[page_idx].extract_text() or ""
            return page_idx + 1, text
    except Exception as e:
        print(f"Error extracting PDF page {page_idx}: {e}")
    return page_idx + 1, ""

def extract_pdf_pages_parallel(content_bytes: bytes) -> tuple[list[tuple[int, str]], int]:
    reader = pypdf.PdfReader(io.BytesIO(content_bytes))
    total_pages = len(reader.pages)
    if total_pages == 0:
        return [], 0
    if total_pages == 1:
        text = reader.pages[0].extract_text() or ""
        return [(1, text)], 1
    
    tasks = [(content_bytes, i) for i in range(total_pages)]
    max_workers = min(8, total_pages)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        pages_data = list(executor.map(_extract_pdf_page, tasks))
    pages_data.sort(key=lambda x: x[0])
    return pages_data, total_pages

def extract_pdf_text_parallel(content_bytes: bytes) -> tuple[str, int]:
    pages_data, total_pages = extract_pdf_pages_parallel(content_bytes)
    full_text = "\n".join(t for _, t in pages_data if t.strip())
    return full_text, total_pages

def document_loader(file_path):
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    
    text = ""
    if file_path.suffix.lower() == ".pdf":
        content_bytes = file_path.read_bytes()
        text, _ = extract_pdf_text_parallel(content_bytes)
    elif file_path.suffix.lower() == ".txt":
        text = file_path.read_text(encoding="utf-8")
    elif file_path.suffix.lower() == ".docx":
        document = docx.Document(str(file_path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    else:
        raise ValueError(f"Unsupported file type: {file_path.suffix}")
    return text

def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 150) -> list[str]:
    if not text:
        return []
    chunks = []
    text = RE_LINE_BREAKS.sub('\n', text)
    start = 0
    text_len = len(text)
    
    while start < text_len:
        end = start + chunk_size
        if end >= text_len:
            chunks.append(text[start:].strip())
            break
            
        split_candidates = [
            text.rfind('\n', start + chunk_size - 100, end),
            text.rfind('. ', start + chunk_size - 100, end),
            text.rfind(' ', start + chunk_size - 50, end)
        ]
        
        split_point = -1
        for candidate in split_candidates:
            if candidate != -1:
                split_point = candidate
                break
                
        if split_point == -1:
            split_point = end
            
        chunks.append(text[start:split_point].strip())
        start = split_point - overlap
        if start < 0:
            start = 0
        if start >= split_point:
            start = split_point + 1
            
    return [c for c in chunks if c.strip()]

def chunk_text_with_pages(pages_data: list[tuple[int, str]], chunk_size: int = 1000, overlap: int = 150) -> tuple[list[str], list[int]]:
    chunks = []
    chunk_pages = []
    for page_num, page_text in pages_data:
        if not page_text or not page_text.strip():
            continue
        page_chunks = chunk_text(page_text, chunk_size=chunk_size, overlap=overlap)
        for c in page_chunks:
            chunks.append(c)
            chunk_pages.append(page_num)
    return chunks, chunk_pages

def load_text_from_document(document):
    """Returns (text, total_pages, pages_data) where pages_data is list of (page_num, page_text)."""
    if isinstance(document, (str, Path)):
        file_path = Path(document)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            content_bytes = file_path.read_bytes()
            pages_data, total_pages = extract_pdf_pages_parallel(content_bytes)
            full_text = "\n".join(t for _, t in pages_data if t.strip())
            return full_text, total_pages, pages_data
        elif suffix == ".txt":
            text = file_path.read_text(encoding="utf-8")
            return text, 1, [(1, text)]
        elif suffix == ".docx":
            doc = docx.Document(str(file_path))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return text, 1, [(1, text)]
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

    if hasattr(document, "read") and hasattr(document, "name"):
        suffix = Path(document.name).suffix.lower()
        content_bytes = document.read()
        if hasattr(document, "seek"):
            document.seek(0)
            
        if suffix == ".pdf":
            pages_data, total_pages = extract_pdf_pages_parallel(content_bytes)
            full_text = "\n".join(t for _, t in pages_data if t.strip())
            return full_text, total_pages, pages_data
        elif suffix == ".txt":
            text = content_bytes.decode("utf-8", errors="ignore")
            return text, 1, [(1, text)]
        elif suffix == ".docx":
            doc = docx.Document(io.BytesIO(content_bytes))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return text, 1, [(1, text)]
        else:
            raise ValueError("Unsupported document input type")
    raise ValueError("Unsupported document input type")

def create_lightweight_embedding(text, dimensions=384):
    tokens = RE_WORD_TOKENS.findall(text.lower())
    if not tokens:
        return [0.0] * dimensions
    vector = [0.0] * dimensions
    for token in tokens:
        index = abs(hash(token)) % dimensions
        vector[index] += 1.0
    norm = sum(value * value for value in vector) ** 0.5
    if norm == 0:
        return [0.0] * dimensions
    return [value / norm for value in vector]

def embedding_generation(chunks):
    token = get_hf_token()
    if token:
        os.environ["HF_TOKEN"] = token
    try:
        embedding_model = get_embedding_model()
        if isinstance(chunks, str):
            chunks = [chunks]
        return embedding_model.encode(chunks, batch_size=32, show_progress_bar=False)
    except Exception as exc:
        print(f"Embedding model failed ({exc}). Using lightweight fallback embeddings.")
        if isinstance(chunks, str):
            chunks = [chunks]
        return [create_lightweight_embedding(chunk) for chunk in chunks]

def vector_store_creation(chunks, embeddings, doc_id="default", file_name="Uploaded Document", chunk_pages=None):
    """
    Upserts chunks for doc_id into the shared 'document_embeddings' collection.
    Cleans previous chunks for this doc_id before inserting to avoid duplicates.
    """
    client = get_chroma_client()
    collection_name = "document_embeddings"
    collection = client.get_or_create_collection(collection_name)
    
    try:
        collection.delete(where={"doc_id": doc_id})
    except Exception:
        pass

    if not chunks:
        return collection

    embedding_values = embeddings.tolist() if hasattr(embeddings, "tolist") else embeddings
    if chunk_pages is None or len(chunk_pages) != len(chunks):
        chunk_pages = [1] * len(chunks)

    ids = [f"{doc_id}_chunk_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "doc_id": doc_id,
            "doc_name": file_name,
            "page": int(chunk_pages[i]),
            "chunk_idx": i
        }
        for i in range(len(chunks))
    ]

    batch_size = 200
    for i in range(0, len(chunks), batch_size):
        end = i + batch_size
        collection.add(
            ids=ids[i:end],
            embeddings=embedding_values[i:end],
            documents=chunks[i:end],
            metadatas=metadatas[i:end]
        )
    return collection

def query_processing(query, vector_store, doc_ids=None):
    query_embedding = embedding_generation(query)
    if isinstance(query_embedding, list) and not isinstance(query_embedding[0], list):
        query_embeddings = [query_embedding]
    else:
        query_embeddings = query_embedding.tolist() if hasattr(query_embedding, "tolist") else query_embedding
        
    where = None
    if doc_ids:
        doc_ids_list = list(doc_ids) if isinstance(doc_ids, (list, tuple, set)) else [doc_ids]
        if len(doc_ids_list) == 1:
            where = {"doc_id": doc_ids_list[0]}
        elif len(doc_ids_list) > 1:
            where = {"doc_id": {"$in": doc_ids_list}}

    try:
        if where:
            matched = vector_store.get(where=where)
            total_chunks = len(matched['ids'])
        else:
            all_ids = vector_store.get()['ids']
            total_chunks = len(all_ids)
    except Exception:
        total_chunks = 10
        
    if total_chunks == 0:
        return {"documents": [[]], "metadatas": [[]], "distances": [[]]}
        
    n_results = total_chunks if total_chunks <= 25 else 10
    n_results = min(n_results, total_chunks)
    if n_results <= 0:
        n_results = 1
        
    return vector_store.query(
        query_embeddings=query_embeddings,
        n_results=n_results,
        where=where
    )

def clean_context_text(text: str) -> str:
    cleaned = RE_MULTIPLE_NEWLINES.sub(" ", text)
    cleaned = RE_MULTIPLE_SPACES.sub(" ", cleaned).strip()
    cleaned = RE_PAGE_NUMBERS.sub("", cleaned)
    cleaned = RE_PP_NUMBERS.sub("", cleaned)
    return cleaned.strip()

def extract_citations_and_context(results: dict) -> tuple[list[dict], str]:
    documents = results.get("documents", [[]])
    metadatas = results.get("metadatas", [[]])
    distances = results.get("distances", [[]])
    
    docs_list = documents[0] if documents and len(documents) > 0 else []
    meta_list = metadatas[0] if metadatas and len(metadatas) > 0 else []
    dist_list = distances[0] if distances and len(distances) > 0 else []
    
    citations = []
    flattened_texts = []
    
    for i, doc_text in enumerate(docs_list):
        if not doc_text:
            continue
        meta = meta_list[i] if i < len(meta_list) and isinstance(meta_list[i], dict) else {}
        dist = dist_list[i] if i < len(dist_list) else None
        score = round(max(0.0, 1.0 - (dist / 2.0)), 4) if dist is not None else 1.0
        
        citations.append({
            "chunk_id": f"chunk_{i}",
            "doc_id": meta.get("doc_id", ""),
            "doc_name": meta.get("doc_name", "Uploaded Document"),
            "page": meta.get("page", 1),
            "text": doc_text,
            "score": score
        })
        flattened_texts.append(doc_text)
        
    context_text = clean_context_text(" ".join(flattened_texts))
    return citations, context_text

def extract_context(results):
    _, context = extract_citations_and_context(results)
    return context

# LLM Providers (Groq -> Ollama -> Heuristic)
def call_groq_api(prompt, temperature=0.0, stream=False):
    api_key = get_groq_api_key()
    if not api_key:
        raise ValueError("GROQ_API_KEY not configured.")
        
    url = "https://api.groq.com/openai/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    
    models_env = os.getenv("GROQ_MODELS")
    if models_env:
        models = [m.strip() for m in models_env.split(",") if m.strip()]
    else:
        models = [
            "llama-3.1-8b-instant",
            "gemma2-9b-it",
            "llama-3.3-70b-versatile",
            "qwen/qwen3.6-27b",
            "openai/gpt-oss-120b",
            "openai/gpt-oss-20b"
        ]
        
    timeout = int(os.getenv("GROQ_TIMEOUT", "12"))
    last_err = None
    
    for model in models:
        try:
            payload = {
                "model": model,
                "messages": [{"role": "user", "content": prompt}],
                "temperature": temperature,
                "max_tokens": 1024,
                "stream": stream
            }
            if stream:
                response = HTTP_SESSION.post(url, headers=headers, json=payload, timeout=timeout, stream=True)
                response.raise_for_status()
                
                def generate_chunks():
                    for line in response.iter_lines():
                        if not line:
                            continue
                        line_str = line.decode('utf-8')
                        if line_str.startswith("data: "):
                            data_str = line_str[6:].strip()
                            if data_str == "[DONE]":
                                break
                            try:
                                data = json.loads(data_str)
                                delta = data["choices"][0]["delta"]
                                content = delta.get("content", "")
                                if content:
                                    yield content
                            except Exception:
                                pass
                return generate_chunks()
            else:
                response = HTTP_SESSION.post(url, headers=headers, json=payload, timeout=timeout)
                response.raise_for_status()
                res_json = response.json()
                return res_json["choices"][0]["message"]["content"].strip()
        except Exception as e:
            last_err = e
            continue
    raise last_err or RuntimeError("All Groq models failed.")

def generate_llm_response(prompt, temperature=0.0, timeout=30, stream=False):
    # Tier 1: Try Groq API
    api_key = get_groq_api_key()
    if api_key:
        try:
            return call_groq_api(prompt, temperature=temperature, stream=stream), "LLM Groq API"
        except Exception as e:
            print(f"Groq API failed: {e}. Falling back to local Ollama.")
            
    # Tier 2: Try local Ollama
    try:
        if stream:
            response = HTTP_SESSION.post(
                "http://localhost:11434/api/generate",
                json={
                    "model": "gemma:2b",
                    "prompt": prompt,
                    "stream": True,
                    "options": {"temperature": temperature}
                },
                timeout=timeout,
                stream=True
            )
            response.raise_for_status()
            def generate_ollama_chunks():
                for line in response.iter_lines():
                    if not line:
                        continue
                    try:
                        data = json.loads(line.decode('utf-8'))
                        chunk = data.get("response", "")
                        if chunk:
                            yield chunk
                    except Exception:
                        pass
            return generate_ollama_chunks(), "Ollama (Local Fallback)"
        else:
            response = HTTP_SESSION.post(
                "http://localhost:11434/api/generate",
                json={
                    "model": "gemma:2b",
                    "prompt": prompt,
                    "stream": False,
                    "options": {"temperature": temperature}
                },
                timeout=timeout
            )
            response.raise_for_status()
            return response.json()["response"].strip(), "Ollama (Local Fallback)"
    except Exception as e:
        raise RuntimeError(f"Ollama local service failed or not reachable: {e}")

# Sentence Segmentation & Translation Helpers
def split_into_sentences(text: str) -> list[str]:
    """
    Split text into sentences respecting abbreviations, decimal numbers, list numbers,
    and Devanagari danda punctuation (। and ॥).
    """
    if not text:
        return []
    sentences = []
    start = 0
    for m in PUNCT_SPLIT.finditer(text):
        punct = m.group(1)
        end_idx = m.end()
        if '.' in punct and start < m.start(1):
            before = text[start:m.start(1)].strip()
            tokens = before.split()
            last_word = tokens[-1].lower() if tokens else ''
            if (last_word + '.') in ABBREVIATIONS or last_word in ABBREVIATIONS:
                continue
            if last_word.isdigit():
                continue
            if len(last_word) == 1 and last_word.isalpha():
                continue
        candidate = text[start:end_idx].strip()
        if candidate:
            sentences.append(candidate)
            start = end_idx
    if start < len(text):
        tail = text[start:].strip()
        if tail:
            sentences.append(tail)
    return sentences

def split_text_on_words(text: str, max_chars: int = 1800) -> list[str]:
    """Split text on word boundaries without dropping or truncating any words."""
    if not text or len(text) <= max_chars:
        return [text] if text else []
    words = text.split(" ")
    chunks = []
    curr = []
    curr_len = 0
    for w in words:
        w_len = len(w)
        if curr_len + w_len + (1 if curr else 0) <= max_chars:
            curr.append(w)
            curr_len += w_len + (1 if len(curr) > 1 else 0)
        else:
            if curr:
                chunks.append(" ".join(curr))
            if w_len > max_chars:
                for i in range(0, w_len, max_chars):
                    chunks.append(w[i:i + max_chars])
                curr = []
                curr_len = 0
            else:
                curr = [w]
                curr_len = w_len
    if curr:
        chunks.append(" ".join(curr))
    return chunks

def translate_with_sarvam(text: str, target_lang_code: str) -> str:
    api_key = get_sarvam_api_key()
    if not api_key:
        raise ValueError("SARVAM_API_KEY not configured.")
        
    url = "https://api.sarvam.ai/translate"
    headers = {
        "api-subscription-key": api_key,
        "Content-Type": "application/json"
    }
    
    # sarvam-translate:v1 max limit is 2000 chars; safe threshold 1800 chars without dropping words
    text_chunks = split_text_on_words(text, max_chars=1800)
    translated_chunks = []
    
    for chunk in text_chunks:
        if not chunk.strip():
            translated_chunks.append(chunk)
            continue
        payload = {
            "input": chunk,
            "source_language_code": "en-IN",
            "target_language_code": target_lang_code,
            "model": "sarvam-translate:v1"
        }
        resp = HTTP_SESSION.post(url, headers=headers, json=payload, timeout=12)
        resp.raise_for_status()
        res_json = resp.json()
        translated_chunks.append(res_json.get("translated_text", chunk))
        
    return " ".join(translated_chunks)

def translate_with_llm(text: str, target_lang_name: str) -> tuple[str, str]:
    prompt = f"""Translate the following text into {target_lang_name}. Preserve meaning, tone, and any factual details exactly. Return only the translated text, no preamble.

Text:
{text}"""
    translated_text, source = generate_llm_response(prompt, temperature=0.1, timeout=30, stream=False)
    return translated_text, source

def translate_sentence_with_fallback(sentence: str, target_lang: str) -> tuple[str, str]:
    """Translate a single sentence through the 3-tier fallback and return (translated_text, tier_name)."""
    if not sentence.strip() or target_lang in ["English", "en"]:
        return sentence, ""
        
    lang_code = LANGUAGE_CODES.get(target_lang)
    if not lang_code:
        return sentence, ""
        
    # Tier 1: Sarvam AI
    if get_sarvam_api_key():
        try:
            translated = translate_with_sarvam(sentence, lang_code)
            return translated, f"Sarvam AI ({target_lang})"
        except Exception as e:
            print(f"Sarvam translation failed: {e}. Falling back to LLM.")

    # Tier 2: LLM Fallback
    try:
        translated, source = translate_with_llm(sentence, target_lang)
        return translated, f"{source} ({target_lang})"
    except Exception as e:
        print(f"LLM translation failed: {e}. Returning original sentence with note.")

    # Tier 3: Return original sentence
    return sentence, f"Translation Unavailable ({target_lang})"

def translate_text(text: str, target_lang: str = "English") -> tuple[str, str | None]:
    if not text or target_lang in ["English", "en"]:
        return text, None
        
    lang_code = LANGUAGE_CODES.get(target_lang)
    if not lang_code:
        return text, None
        
    # Tier 1: Sarvam AI Translate API
    sarvam_key = get_sarvam_api_key()
    if sarvam_key:
        try:
            translated = translate_with_sarvam(text, lang_code)
            return translated, f"Sarvam AI ({target_lang})"
        except Exception as e:
            print(f"Sarvam translation failed: {e}. Falling back to LLM translation.")
    else:
        print("SARVAM_API_KEY not configured. Falling back to LLM translation.")
        
    # Tier 2: LLM Fallback (Groq / Ollama)
    try:
        translated, source = translate_with_llm(text, target_lang)
        return translated, f"{source} ({target_lang})"
    except Exception as e:
        print(f"LLM translation failed: {e}. Returning original text with note.")
        
    # Tier 3: Return original text untranslated with note
    return text, f"Translation Unavailable ({target_lang})"

# Context Retrieval & Heuristic Fallback
def build_fallback_answer(query: str, context: str) -> str:
    if not context:
        return "I don't have enough context to answer that question."
    cleaned_context = clean_context_text(context)
    sentences = split_into_sentences(cleaned_context)
    query_terms = RE_WORD_TOKENS.findall(query.lower())
    
    scored_sentences = []
    for sentence in sentences:
        if len(sentence) < 20:
            continue
        sentence_lower = sentence.lower()
        score = sum(1 for term in query_terms if term in sentence_lower)
        scored_sentences.append((score, sentence.strip()))
        
    scored_sentences.sort(key=lambda item: item[0], reverse=True)
    selected = [s for score, s in scored_sentences if score > 0][:3]
    return " ".join(selected) if selected else cleaned_context[:300]

def context_retrieval(query: str, context: str, ingested_data=None, stream=False):
    metadata_str = ""
    if ingested_data:
        file_name = ingested_data.get("file_name", "Uploaded Document")
        total_pages = ingested_data.get("total_pages", "N/A")
        total_chunks = len(ingested_data.get("chunks", []))
        metadata_str = f"Document Metadata:\n- File Name: {file_name}\n- Total Pages: {total_pages}\n- Total Chunks: {total_chunks}\n\n"

    prompt = f"""Instructions: You are a helpful and factual document assistant. Answer the Question based on the provided Context and Document Metadata.
Be direct and detailed in your answer. If the provided information does not contain the answer, reply with: "I don't know based on the given information."

{metadata_str}Context: {context}

Question: {query}
Answer:"""

    try:
        return generate_llm_response(prompt, temperature=0.0, timeout=30, stream=stream)
    except Exception as e:
        print(f"LLM generation failed ({e}). Attempting internal fallback scoring algorithm.")
        fallback_ans = build_fallback_answer(query, context)
        if stream:
            def single_chunk():
                yield fallback_ans
            return single_chunk(), "Local Fallback (Heuristic Scoring)"
        return fallback_ans, "Local Fallback (Heuristic Scoring)"

# Backward-compatible Result Tuple with Modern Attributes
class AnswerResult(tuple):
    """
    A 2-tuple (text_or_stream, source) for backwards compatibility with:
      ans, source = rag_service.answer_query(...)
    Also exposes:
      result.citations: list of citation dicts
      result.metrics: dict of timing and tier metrics
    """
    def __new__(cls, text_or_stream, source, citations=None, metrics=None):
        instance = super().__new__(cls, (text_or_stream, source))
        instance.citations = citations or []
        instance.metrics = metrics or {}
        return instance

# Ingestion Pipeline
def ingest_document(document, doc_id: str | None = None) -> dict:
    if hasattr(document, "read"):
        content_bytes = document.read()
        if hasattr(document, "seek"):
            document.seek(0)
    elif isinstance(document, (str, Path)):
        content_bytes = Path(document).read_bytes()
    else:
        content_bytes = b""
        
    doc_hash = hashlib.sha256(content_bytes).hexdigest() if content_bytes else "unknown"
    doc_id = doc_id or doc_hash[:12]
    cache_path = CACHE_DIR / f"{doc_hash}.json"
    file_name = getattr(document, "name", Path(document).name if isinstance(document, (str, Path)) else "Uploaded Document")

    # Invalidate query cache for this document ID if replacing
    clear_query_cache(doc_id=doc_id)

    # Check disk cache (verifying schema version 2)
    if cache_path.exists():
        try:
            cached_data = json.loads(cache_path.read_text(encoding="utf-8"))
            if cached_data.get("schema_version") == SCHEMA_VERSION:
                print(f"Loading document chunks and embeddings from disk cache: {cache_path}")
                chunks = cached_data["chunks"]
                chunk_pages = cached_data.get("pages", [1] * len(chunks))
                embeddings = cached_data["embeddings"]
                total_pages = cached_data["total_pages"]
                text = cached_data["text"]

                vector_store_creation(chunks, embeddings, doc_id=doc_id, file_name=file_name, chunk_pages=chunk_pages)
                
                # Update registry
                registry = load_document_registry()
                registry[doc_id] = {
                    "doc_id": doc_id,
                    "doc_hash": doc_hash,
                    "file_name": file_name,
                    "total_pages": total_pages,
                    "chunk_count": len(chunks)
                }
                save_document_registry(registry)

                return {
                    "doc_id": doc_id,
                    "doc_hash": doc_hash,
                    "file_name": file_name,
                    "total_pages": total_pages,
                    "chunks": chunks,
                    "chunk_count": len(chunks),
                    "collection_name": "document_embeddings",
                    "text": text,
                    "cached": True
                }
            else:
                print(f"Cache schema version mismatch for {cache_path}. Treating as cache miss.")
        except Exception as e:
            print(f"Cache load failed: {e}. Rebuilding...")

    # Fresh extraction
    text, total_pages, pages_data = load_text_from_document(document)
    if not text or not text.strip():
        raise ValueError("Unable to extract text from the uploaded document.")

    chunks, chunk_pages = chunk_text_with_pages(pages_data, chunk_size=1000, overlap=150)
    if not chunks:
        chunks = chunk_text(text, chunk_size=1000, overlap=150)
        chunk_pages = [1] * len(chunks)
        
    if not chunks:
        raise ValueError("No text chunks generated from the document.")

    embeddings = embedding_generation(chunks)
    vector_store_creation(chunks, embeddings, doc_id=doc_id, file_name=file_name, chunk_pages=chunk_pages)

    try:
        embedding_list = embeddings.tolist() if hasattr(embeddings, "tolist") else embeddings
        cache_data = {
            "schema_version": SCHEMA_VERSION,
            "doc_hash": doc_hash,
            "file_name": file_name,
            "total_pages": total_pages,
            "chunks": chunks,
            "pages": chunk_pages,
            "embeddings": embedding_list,
            "text": text
        }
        cache_path.write_text(json.dumps(cache_data), encoding="utf-8")
    except Exception as e:
        print(f"Failed to write disk cache: {e}")

    registry = load_document_registry()
    registry[doc_id] = {
        "doc_id": doc_id,
        "doc_hash": doc_hash,
        "file_name": file_name,
        "total_pages": total_pages,
        "chunk_count": len(chunks)
    }
    save_document_registry(registry)

    return {
        "doc_id": doc_id,
        "doc_hash": doc_hash,
        "file_name": file_name,
        "total_pages": total_pages,
        "chunks": chunks,
        "chunk_count": len(chunks),
        "collection_name": "document_embeddings",
        "text": text,
        "cached": False
    }

# Query Execution Pipeline
def answer_query(ingested_data, query: str, target_lang: str = "English", stream: bool = False, doc_ids: list[str] | None = None) -> AnswerResult:
    if not ingested_data and not doc_ids:
        return AnswerResult("Please upload a document first.", "System Info", [], {})

    resolved_doc_ids = []
    if doc_ids:
        resolved_doc_ids = list(doc_ids)
    elif isinstance(ingested_data, dict):
        d_id = ingested_data.get("doc_id", ingested_data.get("doc_hash", "default"))
        resolved_doc_ids = [d_id]
        
    normalized_q = query.strip().lower()
    cache_key = (normalized_q, target_lang, tuple(sorted(resolved_doc_ids)))
    
    cached_result = get_query_cache_entry(cache_key)
    if cached_result is not None:
        return cached_result

    t0_start = time.perf_counter()
    client = get_chroma_client()
    collection = client.get_or_create_collection("document_embeddings")
    
    t0_retrieval = time.perf_counter()
    raw_results = query_processing(query, collection, doc_ids=resolved_doc_ids)
    citations, context = extract_citations_and_context(raw_results)
    retrieval_ms = (time.perf_counter() - t0_retrieval) * 1000

    metrics = {
        "retrieval_ms": round(retrieval_ms, 2),
        "generation_ms": 0.0,
        "first_token_ms": None,
        "translation_ms": 0.0,
        "total_ms": 0.0,
        "llm_provider": "",
        "translation_tiers": [],
        "target_lang": target_lang
    }

    t0_gen = time.perf_counter()

    if not stream:
        res, source = context_retrieval(query, context, ingested_data, stream=False)
        generation_ms = (time.perf_counter() - t0_gen) * 1000
        metrics["generation_ms"] = round(generation_ms, 2)
        metrics["llm_provider"] = source

        if target_lang != "English":
            t0_trans = time.perf_counter()
            translated_res, trans_source = translate_text(res, target_lang)
            translation_ms = (time.perf_counter() - t0_trans) * 1000
            metrics["translation_ms"] = round(translation_ms, 2)
            if trans_source:
                metrics["translation_tiers"] = [trans_source]
                source = f"{source} → {trans_source}"
            final_text = translated_res
        else:
            final_text = res

        metrics["total_ms"] = round((time.perf_counter() - t0_start) * 1000, 2)
        answer_obj = AnswerResult(final_text, source, citations=citations, metrics=metrics)
        set_query_cache_entry(cache_key, answer_obj)
        return answer_obj

    # Streaming mode
    res_stream_or_text, source = context_retrieval(query, context, ingested_data, stream=True)
    metrics["llm_provider"] = source

    if target_lang == "English":
        def generate_english_stream():
            collected_chunks = []
            first_token_recorded = False
            if hasattr(res_stream_or_text, "__iter__") and not isinstance(res_stream_or_text, str):
                for chunk in res_stream_or_text:
                    if not first_token_recorded:
                        metrics["first_token_ms"] = round((time.perf_counter() - t0_gen) * 1000, 2)
                        first_token_recorded = True
                    collected_chunks.append(chunk)
                    yield chunk
            else:
                chunk = str(res_stream_or_text)
                metrics["first_token_ms"] = round((time.perf_counter() - t0_gen) * 1000, 2)
                collected_chunks.append(chunk)
                yield chunk

            metrics["generation_ms"] = round((time.perf_counter() - t0_gen) * 1000, 2)
            metrics["total_ms"] = round((time.perf_counter() - t0_start) * 1000, 2)
            full_text = "".join(collected_chunks)
            set_query_cache_entry(cache_key, AnswerResult(full_text, source, citations=citations, metrics=metrics))

        return AnswerResult(generate_english_stream(), source, citations=citations, metrics=metrics)

    # Multilingual Sentence-Pipelined Streaming for Hindi / Marathi
    def generate_multilingual_stream():
        buffer = ""
        tiers_used = []
        first_token_recorded = False
        t_trans_total = 0.0
        collected_translated = []

        def process_and_yield_sentence(sentence_text):
            nonlocal first_token_recorded, t_trans_total
            t0_t = time.perf_counter()
            trans_sent, tier = translate_sentence_with_fallback(sentence_text, target_lang)
            t_trans_total += (time.perf_counter() - t0_t) * 1000
            if tier:
                tiers_used.append(tier)
            if not first_token_recorded:
                metrics["first_token_ms"] = round((time.perf_counter() - t0_gen) * 1000, 2)
                first_token_recorded = True
            collected_translated.append(trans_sent)
            return trans_sent

        if hasattr(res_stream_or_text, "__iter__") and not isinstance(res_stream_or_text, str):
            for token in res_stream_or_text:
                buffer += token
                sentences = split_into_sentences(buffer)
                if len(sentences) > 1:
                    for s in sentences[:-1]:
                        yield process_and_yield_sentence(s) + " "
                    buffer = sentences[-1]
        else:
            buffer = str(res_stream_or_text)

        if buffer.strip():
            yield process_and_yield_sentence(buffer.strip())

        metrics["generation_ms"] = round((time.perf_counter() - t0_gen) * 1000, 2)
        metrics["translation_ms"] = round(t_trans_total, 2)
        metrics["total_ms"] = round((time.perf_counter() - t0_start) * 1000, 2)
        metrics["translation_tiers"] = list(dict.fromkeys(tiers_used))

        final_source = f"{source} → {', '.join(metrics['translation_tiers'])}" if metrics["translation_tiers"] else source
        full_text = " ".join(collected_translated)
        set_query_cache_entry(cache_key, AnswerResult(full_text, final_source, citations=citations, metrics=metrics))

    return AnswerResult(generate_multilingual_stream(), source, citations=citations, metrics=metrics)

# Document Summarization
def build_fallback_summary(text: str) -> str:
    if not text:
        return "No text to summarize."
    cleaned_text = clean_context_text(text)
    sentences = split_into_sentences(cleaned_text)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 15]
    
    if not sentences:
        return text[:300] + "..."
        
    intro = sentences[:3]
    conclusion = sentences[-2:] if len(sentences) > 5 else []
    
    all_words = RE_LONG_WORDS.findall(cleaned_text.lower())
    from collections import Counter
    word_counts = Counter(all_words)
    top_words = [word for word, count in word_counts.most_common(5)]
    
    body_sentences = []
    if len(sentences) > 5:
        middle_sentences = sentences[3:-2]
        scored = []
        for s in middle_sentences:
            score = sum(1 for word in top_words if word in s.lower())
            scored.append((score, s))
        scored.sort(key=lambda x: x[0], reverse=True)
        body_sentences = [s for score, s in scored[:3] if score > 0]
        
    summary_parts = intro
    if body_sentences:
        summary_parts.append("\n\n**Key Details:**")
        summary_parts.extend(body_sentences)
    if conclusion:
        summary_parts.append("\n\n**Conclusion:**")
        summary_parts.extend(conclusion)
        
    return " ".join(summary_parts)

def summarize_document(ingested_data, target_lang: str = "English") -> tuple[str, str]:
    if not ingested_data or not ingested_data.get("text"):
        return "No document text available to summarize.", "System Info"
        
    doc_id = ingested_data.get("doc_id", ingested_data.get("doc_hash", "default"))
    cache_key = ("__SUMMARY__", target_lang, (doc_id,))
    cached = get_query_cache_entry(cache_key)
    if cached is not None:
        return cached[0], cached[1]

    text = ingested_data["text"]
    max_summary_input_chars = 6000
    if len(text) > max_summary_input_chars:
        input_text = text[:4000] + "\n... [text truncated for summarization] ...\n" + text[-2000:]
    else:
        input_text = text

    prompt = f"""Instructions: Provide a concise, comprehensive summary of the following document. Highlight the main topics, key points, and overall conclusion.
Document:
{input_text}

Summary:"""

    try:
        raw_summary, source = generate_llm_response(prompt, temperature=0.3, timeout=45, stream=False)
    except Exception as e:
        print(f"LLM summarization failed ({e}). Using fallback extractive summary.")
        raw_summary, source = build_fallback_summary(text), "Fallback (Extractive Heuristics)"

    if target_lang != "English":
        translated_summary, trans_source = translate_text(raw_summary, target_lang)
        if trans_source:
            source = f"{source} → {trans_source}"
        res_tuple = (translated_summary, source)
    else:
        res_tuple = (raw_summary, source)
        
    set_query_cache_entry(cache_key, AnswerResult(res_tuple[0], res_tuple[1]))
    return res_tuple

def pipeline(document, query=None):
    ingested = ingest_document(document)
    if query is None:
        return ingested
    return answer_query(ingested, query)
